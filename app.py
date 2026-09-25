"""USMS-GH style SMS platform - FastAPI + SQLite + mock gateway."""
import uuid, secrets, csv, io, time
from datetime import datetime
from fastapi import FastAPI, HTTPException, Query, UploadFile, File, Form, Header
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.gzip import GZipMiddleware
from pydantic import BaseModel, Field
from pathlib import Path

import database as db
import sms_gateway as gw
import auth as authlib

FROM_NUMBER = "+15550001111"
BASE = Path(__file__).parent

app = FastAPI(title="USMS-GH SMS Platform", version="2.0.0")
app.add_middleware(GZipMiddleware, minimum_size=500)

_stats_cache = {"t": 0.0, "v": None}

# ---------- schemas ----------
class ContactIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    phone: str = Field(min_length=1, max_length=20)
    group_name: str = Field(default="General", max_length=60)

class SendIn(BaseModel):
    to: str = Field(min_length=1, max_length=20)
    body: str = Field(min_length=1, max_length=1600)
    sender: str = Field(default="NOVA", max_length=20)

class BulkIn(BaseModel):
    to: list[str] = Field(min_length=1, max_length=500)
    body: str = Field(min_length=1, max_length=1600)
    sender: str = Field(default="NOVA", max_length=20)

class InboundIn(BaseModel):
    from_phone: str = Field(min_length=1, max_length=20, alias="from")
    body: str = Field(min_length=1, max_length=1600)
    class Config:
        populate_by_name = True

class BlackIn(BaseModel):
    phone: str
    reason: str = ""

class SchedIn(BaseModel):
    to: list[str] = Field(min_length=1)
    body: str = Field(min_length=1, max_length=1600)
    send_at: str = "now"
    sender: str = "USMS-GH"

class TicketIn(BaseModel):
    subject: str = Field(min_length=1, max_length=140)
    message: str = Field(min_length=1, max_length=2000)

class KeyIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)

class TopupIn(BaseModel):
    amount: float = Field(gt=0, le=100000)

class RegisterIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    email: str = Field(min_length=3, max_length=120)
    password: str = Field(min_length=6, max_length=128)
    phone: str = Field(default="", max_length=20)

class LoginIn(BaseModel):
    email: str = Field(min_length=3, max_length=120)
    password: str = Field(min_length=1, max_length=128)
    role: str = Field(default="customer", max_length=20)

class SenderIn(BaseModel):
    value: str = Field(min_length=1, max_length=11)

class TemplateIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    body: str = Field(min_length=1, max_length=1600)

class AdminTopupIn(BaseModel):
    amount: float = Field(gt=-100000, lt=100000)

def current_user(authorization: str = Header(default="")) -> dict | None:
    if not authorization.startswith("Bearer "):
        return None
    token = authorization[7:].strip()
    if not token:
        return None
    conn = db.get_conn()
    s = conn.execute("SELECT user_id, expires_at FROM sessions WHERE token=?", (token,)).fetchone()
    if not s:
        conn.close()
        return None
    try:
        if s["expires_at"] < datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"):
            conn.execute("DELETE FROM sessions WHERE token=?", (token,))
            conn.commit(); conn.close()
            return None
    except Exception:
        pass
    u = conn.execute("SELECT id,name,email,phone,role,balance,created_at FROM users WHERE id=?", (s["user_id"],)).fetchone()
    conn.close()
    return db.row_to_dict(u) if u else None

def require_user(authorization: str = Header(default="")) -> dict:
    u = current_user(authorization)
    if not u:
        raise HTTPException(401, "Login required")
    return u

def require_admin(authorization: str = Header(default="")) -> dict:
    u = require_user(authorization)
    if u["role"] != "admin":
        raise HTTPException(403, "Admin only")
    return u

# ---------- helpers ----------
def find_contact_id(phone):
    conn = db.get_conn()
    r = conn.execute("SELECT id FROM contacts WHERE phone=?", (phone,)).fetchone()
    conn.close()
    return r["id"] if r else None

def is_blacklisted(phone):
    conn = db.get_conn()
    r = conn.execute("SELECT 1 FROM blacklist WHERE phone=?", (phone,)).fetchone()
    conn.close()
    return bool(r)

def insert_outbound(to, body, sender="USMS-GH", bulk_id=None):
    to = gw.normalize_phone(to)
    err = gw.validate(to, body)
    if not err and is_blacklisted(to):
        err = "Blocked: number is on blacklist"
    conn = db.get_conn()
    cid = None
    try:
        r = conn.execute("SELECT id FROM contacts WHERE phone=?", (to,)).fetchone()
        cid = r["id"] if r else None
    except Exception:
        pass
    status = "failed" if err else "queued"
    cur = conn.execute(
        "INSERT INTO messages (direction,from_phone,to_phone,body,sender,status,contact_id,bulk_id,error) VALUES ('outbound',?,?,?,?,?,?,?,?)",
        (FROM_NUMBER, to, body, sender, status, cid, bulk_id, err),
    )
    mid = cur.lastrowid
    conn.commit()
    row = conn.execute("SELECT * FROM messages WHERE id=?", (mid,)).fetchone()
    conn.close()
    if not err:
        gw.send_async(mid)
        # deduct 1 credit per segment-ish
        try:
            bal = float(db.get_setting("balance", "150"))
            db.set_setting("balance", str(max(0, bal - 1)))
        except Exception:
            pass
    _stats_cache["t"] = 0
    return db.row_to_dict(row)

@app.on_event("startup")
def _startup():
    db.init_db()

# ---------- meta ----------
@app.get("/api/health")
def health():
    return {"ok": True, "gateway": "mock", "from": FROM_NUMBER}

@app.get("/api/stats")
def stats():
    now = time.time()
    if _stats_cache["v"] and now - _stats_cache["t"] < 2.0:
        return _stats_cache["v"]
    conn = db.get_conn()
    c = lambda q, p=(): conn.execute(q, p).fetchone()[0]
    try:
        groups = c("SELECT COUNT(DISTINCT group_name) FROM contacts")
    except Exception:
        groups = 0
    out = {
        "contacts": c("SELECT COUNT(*) FROM contacts"),
        "groups": groups,
        "outbound": c("SELECT COUNT(*) FROM messages WHERE direction='outbound'"),
        "inbound": c("SELECT COUNT(*) FROM messages WHERE direction='inbound'"),
        "delivered": c("SELECT COUNT(*) FROM messages WHERE status='delivered'"),
        "failed": c("SELECT COUNT(*) FROM messages WHERE status='failed'"),
        "pending": c("SELECT COUNT(*) FROM messages WHERE status IN ('queued','sent')"),
        "blacklist": c("SELECT COUNT(*) FROM blacklist"),
        "scheduled": c("SELECT COUNT(*) FROM scheduled WHERE status='pending'"),
        "balance": float(db.get_setting("balance", "150")),
    }
    conn.close()
    _stats_cache["t"] = time.time()
    _stats_cache["v"] = out
    return out

@app.get("/api/announcements")
def announcements():
    return [{"title": "Exciting New Features Are Here!",
             "body": "Bulk SMS upgraded: Sub-Accounts for teams/resellers and Support Tickets right from your dashboard.",
             "when": "recent"}]

@app.post("/api/topup")
def topup(t: TopupIn):
    bal = float(db.get_setting("balance", "150")) + t.amount
    db.set_setting("balance", str(bal))
    _stats_cache["t"] = 0
    return {"balance": bal}

# ---------- contacts ----------
@app.get("/api/contacts")
def list_contacts(q: str = "", group: str = "", limit: int = Query(300, le=1000)):
    conn = db.get_conn()
    clauses, params = [], []
    if q:
        clauses.append("(name LIKE ? OR phone LIKE ?)"); params += [f"%{q}%", f"%{q}%"]
    if group:
        clauses.append("group_name=?"); params.append(group)
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
    try:
        rows = conn.execute(f"SELECT * FROM contacts {where} ORDER BY id DESC LIMIT ?", (*params, limit)).fetchall()
    except Exception:
        rows = conn.execute("SELECT id,name,phone,created_at FROM contacts ORDER BY id DESC").fetchall()
        rows = [dict(r) | {"group_name": "General"} for r in rows]
        conn.close()
        return rows
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.get("/api/groups")
def groups():
    conn = db.get_conn()
    rows = conn.execute("SELECT id, name, created_at FROM groups ORDER BY name").fetchall()
    out = []
    for r in rows:
        n = conn.execute("SELECT COUNT(*) FROM contacts WHERE group_name=?", (r["name"],)).fetchone()[0]
        d = db.row_to_dict(r); d["count"] = n
        out.append(d)
    conn.close()
    return out

class GroupIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)

@app.post("/api/groups", status_code=201)
def create_group(g: GroupIn, authorization: str = Header(default="")):
    require_user(authorization)
    name = g.name.strip()
    if not name:
        raise HTTPException(400, "Group name required")
    conn = db.get_conn()
    try:
        cur = conn.execute("INSERT INTO groups (name) VALUES (?)", (name,))
        conn.commit()
        row = conn.execute("SELECT id, name, created_at FROM groups WHERE id=?", (cur.lastrowid,)).fetchone()
    except Exception:
        conn.close()
        raise HTTPException(400, "Group already exists")
    conn.close()
    d = db.row_to_dict(row); d["count"] = 0
    return d

@app.delete("/api/groups/{gid}")
def delete_group(gid: int, authorization: str = Header(default="")):
    require_user(authorization)
    conn = db.get_conn()
    r = conn.execute("SELECT name FROM groups WHERE id=?", (gid,)).fetchone()
    if not r:
        conn.close()
        raise HTTPException(404, "Group not found")
    if r["name"] == "General":
        conn.close()
        raise HTTPException(400, "Cannot delete General group")
    n = conn.execute("SELECT COUNT(*) FROM contacts WHERE group_name=?", (r["name"],)).fetchone()[0]
    if n > 0:
        conn.close()
        raise HTTPException(400, f"Group has {n} contact(s). Move or delete them first.")
    conn.execute("DELETE FROM groups WHERE id=?", (gid,))
    conn.commit(); conn.close()
    return {"deleted": gid}

@app.post("/api/contacts", status_code=201)
def create_contact(c: ContactIn, authorization: str = Header(default="")):
    require_user(authorization)
    phone = gw.normalize_phone(c.phone)
    if not gw.PHONE_RE.match(phone):
        raise HTTPException(400, "Invalid phone format")
    if is_blacklisted(phone):
        raise HTTPException(400, "Number is blacklisted")
    gname = (c.group_name or "General").strip() or "General"
    conn = db.get_conn()
    if not conn.execute("SELECT 1 FROM groups WHERE name=?", (gname,)).fetchone():
        conn.close()
        raise HTTPException(400, f"Group '{gname}' does not exist. Create the group first, then assign contacts to it.")
    try:
        cur = conn.execute("INSERT INTO contacts (name,phone,group_name) VALUES (?,?,?)",
                           (c.name.strip(), phone, gname))
        conn.commit()
        row = conn.execute("SELECT * FROM contacts WHERE id=?", (cur.lastrowid,)).fetchone()
    except Exception:
        conn.close()
        raise HTTPException(400, "Contact with that phone already exists")
    conn.close()
    return db.row_to_dict(row)

@app.post("/api/contacts/import")
async def import_contacts_file(
    file: UploadFile = File(None),
    group_name: str = Form("General"),
    authorization: str = Header(default=""),
):
    require_user(authorization)
    gname = (group_name or "General").strip() or "General"
    conn = db.get_conn()
    if not conn.execute("SELECT 1 FROM groups WHERE name=?", (gname,)).fetchone():
        conn.close()
        raise HTTPException(400, f"Group '{gname}' does not exist. Create the group first, then import into it.")
    conn.close()
    if file is None or not file.filename:
        raise HTTPException(400, "No file uploaded. Send .csv or .xlsx with columns name, phone.")
    fname = file.filename.lower()
    raw = await file.read()
    rows: list[tuple[str, str]] = []
    try:
        if fname.endswith(".xlsx") or fname.endswith(".xls"):
            try:
                from openpyxl import load_workbook
            except ImportError:
                raise HTTPException(400, "Excel support needs 'openpyxl'. pip install openpyxl, or upload CSV.")
            import io as _io
            wb = load_workbook(filename=_io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            header = None
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                vals = [(str(c).strip() if c is not None else "") for c in r]
                if i == 0 and any(v.lower() in ("name", "phone", "number", "fullname") for v in vals):
                    header = [v.lower() for v in vals]
                    continue
                if header:
                    d = dict(zip(header, vals))
                    rows.append((d.get("name", "") or d.get("fullname", ""), d.get("phone", "") or d.get("number", "")))
                elif len(vals) >= 2:
                    rows.append((vals[0], vals[1]))
                elif len(vals) == 1 and vals[0]:
                    rows.append(("", vals[0]))
        else:
            text = raw.decode("utf-8-sig", errors="replace")
            import csv as _csv, io as _io
            sample = text[:2048]
            try:
                dialect = _csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except Exception:
                dialect = _csv.excel
            reader = _csv.DictReader(_io.StringIO(text), dialect=dialect)
            if reader.fieldnames and any((f or "").strip().lower() in ("name", "phone", "number", "fullname") for f in reader.fieldnames):
                for r in reader:
                    low = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
                    rows.append((low.get("name", "") or low.get("fullname", ""), low.get("phone", "") or low.get("number", "")))
            else:
                for r in _csv.reader(_io.StringIO(text), dialect):
                    if not r or not any(c.strip() for c in r):
                        continue
                    if len(r) >= 2:
                        rows.append((r[0].strip(), r[1].strip()))
                    else:
                        rows.append(("", r[0].strip()))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Could not parse file: {e}")
    imported, skipped = 0, 0
    errors: list[str] = []
    conn = db.get_conn()
    for idx, (nm, ph) in enumerate(rows[:5000], start=1):
        ph = gw.normalize_phone((ph or "").strip())
        nm = (nm or "").strip() or ph
        if not gw.PHONE_RE.match(ph):
            skipped += 1
            if len(errors) < 10:
                errors.append(f"row {idx}: bad phone '{ph}'")
            continue
        if conn.execute("SELECT 1 FROM blacklist WHERE phone=?", (ph,)).fetchone():
            skipped += 1
            if len(errors) < 10:
                errors.append(f"row {idx}: blacklisted '{ph}'")
            continue
        try:
            conn.execute("INSERT INTO contacts (name,phone,group_name) VALUES (?,?,?)", (nm, ph, gname))
            imported += 1
        except Exception:
            skipped += 1
    conn.commit(); conn.close()
    _stats_cache["t"] = 0
    return {"group": gname, "imported": imported, "skipped": skipped, "errors": errors}

@app.delete("/api/contacts/{cid}")
def delete_contact(cid: int):
    conn = db.get_conn()
    conn.execute("UPDATE messages SET contact_id=NULL WHERE contact_id=?", (cid,))
    cur = conn.execute("DELETE FROM contacts WHERE id=?", (cid,))
    conn.commit(); conn.close()
    if cur.rowcount == 0:
        raise HTTPException(404, "Contact not found")
    return {"deleted": cid}

# ---------- blacklist ----------
@app.get("/api/blacklist")
def list_black():
    conn = db.get_conn()
    rows = conn.execute("SELECT * FROM blacklist ORDER BY id DESC").fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/blacklist", status_code=201)
def add_black(b: BlackIn):
    phone = gw.normalize_phone(b.phone)
    conn = db.get_conn()
    try:
        cur = conn.execute("INSERT INTO blacklist (phone,reason) VALUES (?,?)", (phone, b.reason))
        conn.commit()
        row = conn.execute("SELECT * FROM blacklist WHERE id=?", (cur.lastrowid,)).fetchone()
    except Exception:
        conn.close()
        raise HTTPException(400, "Already blacklisted")
    conn.close()
    return db.row_to_dict(row)

@app.delete("/api/blacklist/{bid}")
def del_black(bid: int):
    conn = db.get_conn()
    cur = conn.execute("DELETE FROM blacklist WHERE id=?", (bid,))
    conn.commit(); conn.close()
    if cur.rowcount == 0:
        raise HTTPException(404, "Not found")
    return {"deleted": bid}

# ---------- messaging ----------
def require_approved_sender(sender: str) -> str:
    v = (sender or "").strip().upper()
    if not v:
        raise HTTPException(400, "Pick an approved sender ID first")
    conn = db.get_conn()
    r = conn.execute("SELECT status FROM sender_ids WHERE value=?", (v,)).fetchone()
    conn.close()
    if not r:
        raise HTTPException(400, f"Unknown sender '{v}'. Request it on the Sender ID page first.")
    if r["status"] != "approved":
        raise HTTPException(400, f"Sender '{v}' is {r['status']} — only approved senders can send.")
    return v

@app.post("/api/send", status_code=201)
def send_sms(m: SendIn, authorization: str = Header(default="")):
    require_user(authorization)
    sender = require_approved_sender(m.sender)
    return insert_outbound(m.to, m.body, sender)

@app.post("/api/bulk", status_code=201)
def send_bulk(b: BulkIn, authorization: str = Header(default="")):
    require_user(authorization)
    sender = require_approved_sender(b.sender)
    bulk_id = "bulk_" + uuid.uuid4().hex[:8]
    results = [insert_outbound(p, b.body, sender, bulk_id) for p in b.to]
    ok = sum(1 for r in results if r["status"] != "failed")
    return {"bulk_id": bulk_id, "total": len(results), "accepted": ok,
            "failed": len(results) - ok, "messages": results}

@app.post("/api/inbound", status_code=201)
def simulate_inbound(m: InboundIn):
    phone = gw.normalize_phone(m.from_phone)
    conn = db.get_conn()
    cid = find_contact_id(phone)
    cur = conn.execute(
        "INSERT INTO messages (direction,from_phone,to_phone,body,status,contact_id) VALUES ('inbound',?,?,?,'received',?)",
        (phone, FROM_NUMBER, m.body, cid))
    mid = cur.lastrowid
    conn.commit()
    row = conn.execute("SELECT * FROM messages WHERE id=?", (mid,)).fetchone()
    conn.close()
    return db.row_to_dict(row)

@app.get("/api/messages")
def list_messages(direction: str = "", status: str = "", search: str = "",
                  limit: int = Query(60, le=200), offset: int = Query(0, ge=0)):
    conn = db.get_conn()
    clauses, params = [], []
    if direction in ("inbound", "outbound"):
        clauses.append("direction=?"); params.append(direction)
    if status:
        clauses.append("status=?"); params.append(status)
    if search:
        clauses.append("(to_phone LIKE ? OR from_phone LIKE ? OR body LIKE ?)")
        params += [f"%{search}%"] * 3
    where = ("WHERE " + " AND ".join(clauses)) if clauses else ""
    rows = conn.execute(f"SELECT * FROM messages {where} ORDER BY id DESC LIMIT ? OFFSET ?", (*params, limit, offset)).fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.get("/api/messages/{mid}")
def get_message(mid: int):
    conn = db.get_conn()
    row = conn.execute("SELECT * FROM messages WHERE id=?", (mid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Message not found")
    return db.row_to_dict(row)

@app.get("/api/export.csv")
def export_csv():
    conn = db.get_conn()
    rows = conn.execute("SELECT * FROM messages ORDER BY id DESC LIMIT 2000").fetchall()
    conn.close()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["id", "direction", "from", "to", "body", "sender", "status", "error", "created_at"])
    for r in rows:
        d = dict(r)
        w.writerow([d.get("id"), d.get("direction"), d.get("from_phone"), d.get("to_phone"),
                    d.get("body"), d.get("sender"), d.get("status"), d.get("error"), d.get("created_at")])
    buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=sms-report.csv"})

# ---------- conversations / chat ----------
@app.get("/api/conversations")
def conversations():
    conn = db.get_conn()
    rows = conn.execute("""
      SELECT CASE WHEN direction='outbound' THEN to_phone ELSE from_phone END AS peer,
             COUNT(*) n, MAX(created_at) last,
             SUM(CASE WHEN direction='inbound' THEN 1 ELSE 0 END) inbound_n
      FROM messages GROUP BY peer ORDER BY last DESC LIMIT 100""").fetchall()
    out = []
    for r in rows:
        d = dict(r)
        c = conn.execute("SELECT name FROM contacts WHERE phone=?", (d["peer"],)).fetchone()
        d["name"] = c["name"] if c else None
        last = conn.execute(
            "SELECT * FROM messages WHERE to_phone=? OR from_phone=? ORDER BY id DESC LIMIT 1",
            (d["peer"], d["peer"])).fetchone()
        d["preview"] = last["body"][:80] if last else ""
        out.append(d)
    conn.close()
    return out

@app.get("/api/conversations/{phone}")
def conversation(phone: str):
    conn = db.get_conn()
    rows = conn.execute(
        "SELECT * FROM messages WHERE to_phone=? OR from_phone=? ORDER BY id ASC LIMIT 300",
        (phone, phone)).fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

# ---------- scheduled / automate ----------
@app.get("/api/scheduled")
def list_sched():
    conn = db.get_conn()
    rows = conn.execute("SELECT * FROM scheduled ORDER BY id DESC").fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/scheduled", status_code=201)
def create_sched(s: SchedIn, authorization: str = Header(default="")):
    require_user(authorization)
    sender = require_approved_sender(s.sender)
    conn = db.get_conn()
    cur = conn.execute("INSERT INTO scheduled (to_phones,body,sender,send_at) VALUES (?,?,?,?)",
                       (",".join(s.to), s.body, sender, s.send_at))
    conn.commit()
    row = conn.execute("SELECT * FROM scheduled WHERE id=?", (cur.lastrowid,)).fetchone()
    conn.close()
    d = db.row_to_dict(row)
    if s.send_at.strip().lower() == "now":
        for p in s.to:
            insert_outbound(p, s.body, sender)
        conn = db.get_conn()
        conn.execute("UPDATE scheduled SET status='sent' WHERE id=?", (d["id"],))
        conn.commit(); conn.close()
        d["status"] = "sent"
    return d

@app.delete("/api/scheduled/{sid}")
def del_sched(sid: int):
    conn = db.get_conn()
    cur = conn.execute("DELETE FROM scheduled WHERE id=?", (sid,))
    conn.commit(); conn.close()
    if cur.rowcount == 0:
        raise HTTPException(404, "Not found")
    return {"deleted": sid}

# ---------- tickets / keys ----------
@app.get("/api/tickets")
def list_tickets():
    conn = db.get_conn()
    rows = conn.execute("SELECT * FROM tickets ORDER BY id DESC").fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/tickets", status_code=201)
def create_ticket(t: TicketIn):
    conn = db.get_conn()
    cur = conn.execute("INSERT INTO tickets (subject,message) VALUES (?,?)", (t.subject, t.message))
    conn.commit()
    row = conn.execute("SELECT * FROM tickets WHERE id=?", (cur.lastrowid,)).fetchone()
    conn.close()
    return db.row_to_dict(row)

@app.get("/api/keys")
def list_keys():
    conn = db.get_conn()
    rows = conn.execute("SELECT id,name,key,created_at FROM api_keys ORDER BY id DESC").fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/keys", status_code=201)
def create_key(k: KeyIn):
    key = "usms_" + secrets.token_hex(16)
    conn = db.get_conn()
    cur = conn.execute("INSERT INTO api_keys (name,key) VALUES (?,?)", (k.name, key))
    conn.commit()
    row = conn.execute("SELECT id,name,key,created_at FROM api_keys WHERE id=?", (cur.lastrowid,)).fetchone()
    conn.close()
    return db.row_to_dict(row)

@app.delete("/api/keys/{kid}")
def del_key(kid: int):
    conn = db.get_conn()
    cur = conn.execute("DELETE FROM api_keys WHERE id=?", (kid,))
    conn.commit(); conn.close()
    if cur.rowcount == 0:
        raise HTTPException(404, "Not found")
    return {"deleted": kid}

# ---------- frontend ----------
STATIC = BASE / "static"
app.mount("/static", StaticFiles(directory=STATIC), name="static")

@app.get("/login")
def login_page():
    return FileResponse(STATIC / "login.html", headers={"Cache-Control": "no-store"})

@app.get("/admin")
def admin_page():
    return FileResponse(STATIC / "admin-login.html", headers={"Cache-Control": "no-store"})

@app.get("/")
def index():
    return FileResponse(STATIC / "index.html", headers={"Cache-Control": "no-store"})

# ---------- auth ----------
@app.post("/api/auth/register", status_code=201)
def register(r: RegisterIn):
    email = r.email.strip().lower()
    if "@" not in email:
        raise HTTPException(400, "Invalid email")
    conn = db.get_conn()
    if conn.execute("SELECT 1 FROM users WHERE email=?", (email,)).fetchone():
        conn.close()
        raise HTTPException(400, "Email already registered")
    cur = conn.execute(
        "INSERT INTO users (name,email,phone,password_hash,role,balance) VALUES (?,?,?,?, 'customer',150)",
        (r.name.strip(), email, r.phone.strip(), authlib.hash_password(r.password)))
    conn.commit()
    uid = cur.lastrowid
    token = authlib.new_token()
    conn.execute("INSERT INTO sessions (token,user_id,expires_at) VALUES (?,?,?)",
                 (token, uid, authlib.expiry()))
    conn.commit()
    u = conn.execute("SELECT id,name,email,phone,role,balance,created_at FROM users WHERE id=?", (uid,)).fetchone()
    conn.close()
    return {"token": token, "user": db.row_to_dict(u)}

@app.post("/api/auth/login")
def login(r: LoginIn):
    email = r.email.strip().lower()
    conn = db.get_conn()
    u = conn.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
    conn.close()
    if not u or not authlib.verify_password(r.password, u["password_hash"]):
        raise HTTPException(401, "Invalid email or password")
    want = (r.role or "customer").strip().lower()
    if want in ("admin", "customer") and u["role"] != want:
        raise HTTPException(403, f"This account is not a {want} account. Use the {u['role']} login.")
    token = authlib.new_token()
    conn = db.get_conn()
    conn.execute("INSERT INTO sessions (token,user_id,expires_at) VALUES (?,?,?)",
                 (token, u["id"], authlib.expiry()))
    conn.commit()
    conn.close()
    d = dict(u)
    d.pop("password_hash", None)
    return {"token": token, "user": {k: d[k] for k in ("id", "name", "email", "phone", "role", "balance", "created_at")}}

@app.get("/api/auth/me")
def me(authorization: str = Header(default="")):
    u = require_user(authorization)
    return u

@app.post("/api/auth/logout")
def logout(authorization: str = Header(default="")):
    if authorization.startswith("Bearer "):
        conn = db.get_conn()
        conn.execute("DELETE FROM sessions WHERE token=?", (authorization[7:].strip(),))
        conn.commit(); conn.close()
    return {"ok": True}

class ProfileIn(BaseModel):
    name: str = Field(default="", max_length=100)
    email: str = Field(default="", max_length=120)
    phone: str = Field(default="", max_length=20)

class PasswordIn(BaseModel):
    current: str = Field(min_length=1, max_length=128)
    new: str = Field(min_length=6, max_length=128)

@app.patch("/api/auth/profile")
def update_profile(p: ProfileIn, authorization: str = Header(default="")):
    me = require_user(authorization)
    conn = db.get_conn()
    u = conn.execute("SELECT * FROM users WHERE id=?", (me["id"],)).fetchone()
    name = p.name.strip() or u["name"]
    email = (p.email.strip().lower() or u["email"])
    phone = p.phone.strip() if p.phone != "" else (u["phone"] or "")
    if "@" not in email:
        conn.close()
        raise HTTPException(400, "Invalid email")
    if email != u["email"] and conn.execute("SELECT 1 FROM users WHERE email=?", (email,)).fetchone():
        conn.close()
        raise HTTPException(400, "Email already in use")
    conn.execute("UPDATE users SET name=?, email=?, phone=? WHERE id=?", (name, email, phone, me["id"]))
    conn.commit()
    row = conn.execute("SELECT id,name,email,phone,role,balance,created_at FROM users WHERE id=?", (me["id"],)).fetchone()
    conn.close()
    return db.row_to_dict(row)

@app.post("/api/auth/password")
def change_password(p: PasswordIn, authorization: str = Header(default="")):
    me = require_user(authorization)
    conn = db.get_conn()
    u = conn.execute("SELECT password_hash FROM users WHERE id=?", (me["id"],)).fetchone()
    if not authlib.verify_password(p.current, u["password_hash"]):
        conn.close()
        raise HTTPException(401, "Current password is wrong")
    conn.execute("UPDATE users SET password_hash=? WHERE id=?",
                 (authlib.hash_password(p.new), me["id"]))
    conn.execute("DELETE FROM sessions WHERE user_id=?", (me["id"],))
    token = authlib.new_token()
    conn.execute("INSERT INTO sessions (token,user_id,expires_at) VALUES (?,?,?)",
                 (token, me["id"], authlib.expiry()))
    conn.commit(); conn.close()
    return {"ok": True, "token": token}

@app.get("/api/billing")
def billing(authorization: str = Header(default="")):
    me = require_user(authorization)
    conn = db.get_conn()
    out = {
        "plan": "Free",
        "price": "GH₵0",
        "balance": float(db.get_setting("balance", "150")),
        "outbound": conn.execute("SELECT COUNT(*) FROM messages WHERE direction='outbound'").fetchone()[0],
        "delivered": conn.execute("SELECT COUNT(*) FROM messages WHERE status='delivered'").fetchone()[0],
        "user": me["email"],
    }
    conn.close()
    return out

# ---------- sender IDs (customers request, only admins approve) ----------
@app.get("/api/senders")
def list_senders(authorization: str = Header(default="")):
    me = require_user(authorization)
    conn = db.get_conn()
    if me["role"] == "admin":
        rows = conn.execute("SELECT s.*, u.email AS owner FROM sender_ids s LEFT JOIN users u ON u.id=s.user_id ORDER BY s.id DESC").fetchall()
    else:
        rows = conn.execute(
            "SELECT s.*, u.email AS owner FROM sender_ids s LEFT JOIN users u ON u.id=s.user_id WHERE s.status='approved' OR s.user_id=? ORDER BY s.id DESC",
            (me["id"],)).fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/senders", status_code=201)
def add_sender(s: SenderIn, authorization: str = Header(default="")):
    me = require_user(authorization)
    v = s.value.strip().upper()
    if not (3 <= len(v) <= 11):
        raise HTTPException(400, "Sender ID must be 3-11 chars")
    conn = db.get_conn()
    try:
        cur = conn.execute("INSERT INTO sender_ids (value, user_id) VALUES (?,?)", (v, me["id"]))
        conn.commit()
        row = conn.execute("SELECT s.*, u.email AS owner FROM sender_ids s LEFT JOIN users u ON u.id=s.user_id WHERE s.id=?", (cur.lastrowid,)).fetchone()
    except Exception:
        conn.close()
        raise HTTPException(400, "Sender ID already exists")
    conn.close()
    return db.row_to_dict(row)

@app.post("/api/senders/{sid}/approve")
def approve_sender(sid: int, authorization: str = Header(default="")):
    require_admin(authorization)
    conn = db.get_conn()
    conn.execute("UPDATE sender_ids SET status='approved' WHERE id=?", (sid,))
    conn.commit()
    row = conn.execute("SELECT s.*, u.email AS owner FROM sender_ids s LEFT JOIN users u ON u.id=s.user_id WHERE s.id=?", (sid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Not found")
    return db.row_to_dict(row)

@app.post("/api/senders/{sid}/reject")
def reject_sender(sid: int, authorization: str = Header(default="")):
    require_admin(authorization)
    conn = db.get_conn()
    conn.execute("UPDATE sender_ids SET status='rejected' WHERE id=?", (sid,))
    conn.commit()
    row = conn.execute("SELECT s.*, u.email AS owner FROM sender_ids s LEFT JOIN users u ON u.id=s.user_id WHERE s.id=?", (sid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Not found")
    return db.row_to_dict(row)

@app.delete("/api/senders/{sid}")
def delete_sender(sid: int, authorization: str = Header(default="")):
    me = require_user(authorization)
    conn = db.get_conn()
    r = conn.execute("SELECT user_id, status FROM sender_ids WHERE id=?", (sid,)).fetchone()
    if not r:
        conn.close()
        raise HTTPException(404, "Not found")
    if me["role"] != "admin" and r["user_id"] != me["id"]:
        conn.close()
        raise HTTPException(403, "You can only delete your own sender requests")
    if me["role"] != "admin" and r["status"] == "approved":
        conn.close()
        raise HTTPException(403, "Approved senders can only be removed by an admin")
    conn.execute("DELETE FROM sender_ids WHERE id=?", (sid,))
    conn.commit(); conn.close()
    return {"deleted": sid}

# ---------- templates ----------
@app.get("/api/templates")
def list_templates():
    conn = db.get_conn()
    rows = conn.execute("SELECT * FROM templates ORDER BY id DESC").fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/templates", status_code=201)
def add_template(t: TemplateIn):
    conn = db.get_conn()
    cur = conn.execute("INSERT INTO templates (name,body) VALUES (?,?)", (t.name.strip(), t.body))
    conn.commit()
    row = conn.execute("SELECT * FROM templates WHERE id=?", (cur.lastrowid,)).fetchone()
    conn.close()
    return db.row_to_dict(row)

@app.delete("/api/templates/{tid}")
def delete_template(tid: int):
    conn = db.get_conn()
    cur = conn.execute("DELETE FROM templates WHERE id=?", (tid,))
    conn.commit(); conn.close()
    if cur.rowcount == 0:
        raise HTTPException(404, "Not found")
    return {"deleted": tid}

# ---------- admin ----------
@app.get("/api/admin/users")
def admin_users(authorization: str = Header(default="")):
    require_admin(authorization)
    conn = db.get_conn()
    rows = conn.execute("SELECT id,name,email,phone,role,balance,created_at FROM users ORDER BY id DESC").fetchall()
    conn.close()
    return [db.row_to_dict(r) for r in rows]

@app.post("/api/admin/users/{uid}/topup")
def admin_topup(uid: int, t: AdminTopupIn, authorization: str = Header(default="")):
    require_admin(authorization)
    conn = db.get_conn()
    u = conn.execute("SELECT balance FROM users WHERE id=?", (uid,)).fetchone()
    if not u:
        conn.close()
        raise HTTPException(404, "User not found")
    bal = (u["balance"] or 0) + t.amount
    conn.execute("UPDATE users SET balance=? WHERE id=?", (bal, uid))
    conn.commit(); conn.close()
    return {"id": uid, "balance": bal}

@app.get("/api/activity")
def activity(days: int = Query(7, ge=1, le=31)):
    conn = db.get_conn()
    rows = conn.execute(
        """SELECT date(created_at) d,
           SUM(CASE WHEN direction='outbound' THEN 1 ELSE 0 END) outbound,
           SUM(CASE WHEN status='delivered' THEN 1 ELSE 0 END) delivered,
           SUM(CASE WHEN direction='inbound' THEN 1 ELSE 0 END) inbound,
           SUM(CASE WHEN status='failed' THEN 1 ELSE 0 END) failed
           FROM messages WHERE date(created_at) >= date('now', ?)
           GROUP BY d ORDER BY d""", (f"-{days-1} days",)).fetchall()
    by = {r["d"]: dict(r) for r in rows}
    out = []
    from datetime import timedelta as _td, date as _date
    for i in range(days):
        day = (_date.today() - _td(days=days - 1 - i)).isoformat()
        r = by.get(day, {})
        out.append({"day": day[5:], "outbound": r.get("outbound", 0) or 0,
                    "delivered": r.get("delivered", 0) or 0,
                    "inbound": r.get("inbound", 0) or 0, "failed": r.get("failed", 0) or 0})
    conn.close()
    return out

@app.get("/api/admin/overview")
def admin_overview(authorization: str = Header(default="")):
    require_admin(authorization)
    conn = db.get_conn()
    c = lambda q, p=(): conn.execute(q, p).fetchone()[0]
    out = {
        "users": c("SELECT COUNT(*) FROM users"),
        "customers": c("SELECT COUNT(*) FROM users WHERE role='customer'"),
        "contacts": c("SELECT COUNT(*) FROM contacts"),
        "groups": c("SELECT COUNT(*) FROM groups"),
        "outbound": c("SELECT COUNT(*) FROM messages WHERE direction='outbound'"),
        "inbound": c("SELECT COUNT(*) FROM messages WHERE direction='inbound'"),
        "delivered": c("SELECT COUNT(*) FROM messages WHERE status='delivered'"),
        "failed": c("SELECT COUNT(*) FROM messages WHERE status='failed'"),
        "pending_senders": c("SELECT COUNT(*) FROM sender_ids WHERE status='pending'"),
        "open_tickets": c("SELECT COUNT(*) FROM tickets WHERE status='open'"),
        "blocked": c("SELECT COUNT(*) FROM blacklist"),
    }
    recent = conn.execute(
        "SELECT id,name,email,role,balance,created_at FROM users ORDER BY id DESC LIMIT 5").fetchall()
    pending = conn.execute(
        "SELECT s.*, u.email AS owner FROM sender_ids s LEFT JOIN users u ON u.id=s.user_id WHERE s.status='pending' ORDER BY s.id DESC LIMIT 8").fetchall()
    conn.close()
    out["recent_users"] = [db.row_to_dict(r) for r in recent]
    out["pending_sender_list"] = [db.row_to_dict(r) for r in pending]
    return out
